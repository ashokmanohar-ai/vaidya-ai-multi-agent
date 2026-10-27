"""Deterministic report renderer. Does not diagnose or validate medical correctness."""
from io import BytesIO
from docx import Document
from docx.shared import Inches,Pt
from PIL import Image


def render_draft(reference: str, doctor_name: str, findings: list, instructions: list, images: list[bytes]|None=None,approved_signature_block=False,review_flags=None) -> bytes:
    doc=Document(); sec=doc.sections[0]
    sec.header.paragraphs[0].text='VAIDYA AI | PATHOLOGY DOCUMENTATION — REVIEW REQUIRED'
    doc.add_heading('HISTOPATHOLOGY DRAFT',0)
    doc.add_paragraph('Patient reference: '+reference)
    doc.add_paragraph('Dictated by: '+doctor_name)
    doc.add_paragraph('Not a diagnosis by AI. Doctor review and approval are mandatory.')
    doc.add_heading('Dictated findings',level=1)
    if not findings:doc.add_paragraph('[No supported findings — doctor must clarify]')
    for f in findings[:60]:
        line=f.get('text','') if isinstance(f,dict) else ''
        if line:
            para=doc.add_paragraph()
            para.add_run(line[:1200])
    if images:
        doc.add_heading('Attached pathology images (not diagnostically analyzed)',level=1)
        for raw in images[:5]:
            try:
                with Image.open(BytesIO(raw)) as picture:
                    if picture.format not in ('PNG','JPEG'):raise ValueError('Invalid image type')
                    picture.verify()
                doc.add_picture(BytesIO(raw),width=Inches(4.0))
            except Exception as exc:raise ValueError('Invalid pathology image') from exc
    if review_flags:
        doc.add_heading('Quality flags — clinician review required',level=1)
        for flag in review_flags[:30]:doc.add_paragraph(str(flag)[:350])
    if approved_signature_block:
        doc.add_paragraph('Doctor signature: ____________________')
    doc.add_paragraph('DRAFT — DO NOT RELEASE UNTIL CLINICIAN APPROVES')
    b=BytesIO();doc.save(b);return b.getvalue()


def render_final(reference: str, doctor_name: str, final_text: str) -> bytes:
    if not 1<=len(final_text)<=25000:raise ValueError('Invalid final report length')
    doc=Document()
    doc.add_heading('HISTOPATHOLOGY — DOCTOR APPROVED',0)
    doc.add_paragraph('Patient reference: '+reference)
    doc.add_paragraph('Approved by: '+doctor_name)
    for line in final_text.splitlines(): doc.add_paragraph(line[:2500])
    b=BytesIO();doc.save(b);return b.getvalue()


def extract_text(content: bytes) -> str:
    doc=Document(BytesIO(content))
    return '\n'.join(p.text for p in doc.paragraphs)
