"""Regression tests for clinician-controlled structured review. Synthetic inputs ONLY."""
from pathlib import Path
from docx import Document
from vaidya.structured_review import (FIELDS,candidates,validate_fields,formatted_report,
                                       reconcile_minor_name_spelling)
from vaidya.clinical_review import assess_report, final_document_issues
from vaidya.reports import structured_final_docx
from vaidya.database import Database

PATIENT={'name':'Sample Ananya','code':'DEMO-2007','age':25,'sex':'Female'}
FIELDS_OK={
 'specimen':'Appendix', 'referring_doctors':'Dr. Example One / Dr. Example Two',
 'biopsy_id':'B/1234/2026', 'received_date':'01/06/2026', 'reported_date':'04/06/2026',
 'gross':'Appendix measuring 5 cm, lumen patent, tip normal.',
 'microscopy':'Appendicular mucosa with inflammatory cells. No atypical cells seen.',
 'impression':'Features of acute appendicitis.',
}

def form(**values):
    out={**FIELDS_OK,'attest':'yes',**{'confirmed_'+key:'yes' for key,_,_ in FIELDS},
         **{key:'yes' for key in ('verified_audio','verified_identity','verified_numbers','verified_negations','verified_instructions')}}
    out.update(values)
    return out

def assessment(source='Mrs. Sample Ananya age 25 female specimen appendix gross 5 cm microscopy no atypical cells impression acute appendicitis',patient=PATIENT,audio=True):
    return assess_report(patient,source,[{'text':source}],[],audio_present=audio)

def test_candidates_do_not_invent_dates_or_negation():
    a=assessment('Mrs. Sample Ananya age 25 female specimen appendix received 36 reported 47 '
                 'gross appendix 5 cm microscopy and atypical cells impression appendicitis')
    d={x['key']:x['candidate'] for x in candidates(a)}
    assert d['received_date']=='' and d['reported_date']==''
    assert 'and atypical cells' in d['microscopy']
    assert 'No atypical cells' not in d['microscopy']

def test_all_fields_are_required_and_individually_verified():
    a=assessment()
    out,issues=validate_fields(form(),assessment=a,patient=PATIENT)
    assert issues==[] and out==FIELDS_OK
    _,issues=validate_fields(form(confirmed_microscopy=''),assessment=a,patient=PATIENT)
    assert any('Microscopy: independent' in x for x in issues)
    _,issues=validate_fields(form(microscopy=''),assessment=a,patient=PATIENT)
    assert any('Microscopy: required' in x for x in issues)

def test_no_attestation_no_approval():
    _,issues=validate_fields(form(attest='',verified_audio=''),assessment=assessment(),patient=PATIENT)
    assert any('attestation' in x for x in issues)
    assert any('verified_audio' in x for x in issues)

def test_date_must_exist_and_be_dd_mm_yyyy():
    a=assessment()
    for wrong in ('31/02/2026','01/06/26','2026-06-01','30/13/2026'):
        _,issues=validate_fields(form(received_date=wrong),assessment=a,patient=PATIENT)
        assert any('received_date' in x for x in issues)
    _,issues=validate_fields(form(received_date='05/06/2026',reported_date='04/06/2026'),assessment=a,patient=PATIENT)
    assert any('precede' in x for x in issues)

def test_conflicting_patient_age_never_bypassed():
    a=assessment('Mrs. Sample Ananya age 45 female specimen appendix gross 5 cm')
    corrected,note=reconcile_minor_name_spelling(a,PATIENT,form(resolve_name_spelling='yes',name_resolution_note='Checked age with recording and registration.'))
    assert note=='' and corrected['blocking_issues']

def test_minor_name_asr_spelling_can_be_explicitly_reconciled():
    patient={'name':'Ananya','code':'DEMO-2010','age':25,'sex':'Female'}
    a=assessment('Mrs. Anaya age 25 female gross appendix microscopy no atypical cells',patient)
    assert a['blocking_issues'] and len(a['blocking_issues'])==1
    corrected,note=reconcile_minor_name_spelling(a,patient,form(resolve_name_spelling='yes',name_resolution_note='Listened to source and verified synthetic registration spelling.'))
    assert corrected['blocking_issues']==[] and note
    _,issues=validate_fields(form(),assessment=corrected,patient=patient)
    assert issues==[]
    assert a['blocking_issues']  # raw source warning must remain immutable

def test_different_name_is_not_minor_correction():
    a=assessment('Mrs. Unknown Person age 25 female gross appendix')
    corrected,note=reconcile_minor_name_spelling(a,PATIENT,form(resolve_name_spelling='yes',name_resolution_note='Verified recording for the synthetic patient name.'))
    assert corrected['blocking_issues'] and not note

def test_unsafe_cell_findings_fail_even_after_box_check():
    _,issues=validate_fields(form(microscopy='Sections show neutrophils and atypical cells.'),assessment=assessment(),patient=PATIENT)
    assert any('ambiguous positive' in x for x in issues)

def test_review_does_not_override_mismatched_selected_patient():
    other={'name':'Patient Alpha','code':'DEMO-101','age':45,'sex':'Female'}
    a=assessment('Mrs. Sample Ananya age 25 female specimen appendix',other)
    _,issues=validate_fields(form(),assessment=a,patient=other)
    assert any('PATIENT IDENTITY' in x for x in issues)
    assert any('PATIENT AGE' in x for x in issues)

def test_final_document_is_structured_and_has_source_audit(tmp_path):
    a=assessment();structured,issues=validate_fields(form(),assessment=a,patient=PATIENT)
    assert not issues
    txt=formatted_report(PATIENT,structured)
    assert not final_document_issues(txt,a)
    out=tmp_path/'synthetic-reviewed.docx'
    structured_final_docx(structured, {'id':'fake-1234'}, PATIENT,
                          {'logo_text':'SYNTHETIC PATHOLOGY','address':'LOCAL DEMO ONLY'},out)
    doc=Document(out)
    contents='\n'.join(x.text for x in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    for expected in ('Sample Ananya','B/1234/2026','Dr. Example One','No atypical cells seen','acute appendicitis'):
        assert expected in contents
    assert 'NOT FOR CLINICAL USE' in contents

def test_additive_migration_preserves_old_jobs(tmp_path):
    db=Database(tmp_path/'local.sqlite3');db.init();db.init()
    assert 'review_json' in {x['name'] for x in db.query('PRAGMA table_info(jobs)')}


def test_new_flask_routes_are_present_and_guarded():
    source=(Path(__file__).resolve().parents[1]/'app.py').read_text()
    assert "@app.post('/patients/create')" in source
    assert "@app.post('/reports/<job_id>/approve-structured')" in source
    assert "if request.form.get('synthetic_only')!='yes'" in source
    assert 'if errors:abort(409' in source
    assert "WHERE id=? AND status=?" in source
