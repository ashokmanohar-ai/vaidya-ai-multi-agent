from docx import Document
from vaidya.database import Database,utcnow
from vaidya.reports import render_docx
from vaidya.providers import BedrockConverseProvider


def test_new_skill_version_changes_signature_template(tmp_path):
    db=Database(tmp_path/'data.sqlite');db.seed_demo()
    patient=db.query('SELECT * FROM patients WHERE id=1',one=True)
    hospital=db.query('SELECT * FROM hospitals WHERE id=1',one=True)
    args=({'id':'job-synthetic'},patient,hospital,[{'text':'Synthetic appendix section'}],[],[],[],)
    first=tmp_path/'before.docx'
    render_docx(*args,first,approved_skills='Baseline: dictated content only')
    assert 'Clinician signature:' not in '\n'.join(p.text for p in Document(first).paragraphs)
    upgraded=tmp_path/'after.docx'
    render_docx(*args,upgraded,approved_skills='Baseline\nAPPROVED: Add signature block')
    assert 'Clinician signature:' in '\n'.join(p.text for p in Document(upgraded).paragraphs)


def test_existing_database_adds_skills_version(tmp_path):
    db=Database(tmp_path/'data.sqlite');db.seed_demo()
    db.init()
    columns=[r['name'] for r in db.query('PRAGMA table_info(jobs)')]
    assert 'skills_version' in columns


def test_bedrock_prompt_contains_approved_preferences(monkeypatch):
    monkeypatch.delenv('BEDROCK_MODELS',raising=False)
    provider=BedrockConverseProvider(model='example.synthetic.model',client=object())
    body=provider._request('Synthetic findings','APPROVED: Show signature block')
    assert 'Show signature block' in body['system'][0]['text']
