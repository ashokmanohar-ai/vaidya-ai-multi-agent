"""Audio-first workflow acceptance tests. No live Sarvam or actual patient data."""
from types import SimpleNamespace
from pathlib import Path
import pytest
from jinja2 import Environment, FileSystemLoader
from docx import Document

from vaidya.database import Database
from vaidya.audio_patient import attach_patient_candidate
from vaidya.clinical_review import dictated_patient
from vaidya.structured_review import verified_audio_patient
from vaidya.reports import render_docx, structured_final_docx


def db_without_patients(tmp_path):
    db=Database(tmp_path/'candidate.sqlite3')
    db.seed_demo(include_patients=False)
    assert db.query('SELECT * FROM patients')==[]
    return db


def test_patient_is_created_only_after_audio_response(tmp_path):
    db=db_without_patients(tmp_path)
    with pytest.raises(ValueError,match='nonempty audio-derived'):
        attach_patient_candidate(db,job_id='123',hospital_id=1,transcript='',audio_hashes=['sha'])
    with pytest.raises(ValueError,match='nonempty audio-derived'):
        attach_patient_candidate(db,job_id='123',hospital_id=1,transcript='Mrs. Anaya age 25 female',audio_hashes=[])
    assert db.query('SELECT * FROM patients')==[]
    source='Patient name Mrs. Anaya age 25 female, appendix specimen'
    pid=attach_patient_candidate(db,job_id='123',hospital_id=1,transcript=source,audio_hashes=['sha'])
    patient=db.query('SELECT * FROM patients WHERE id=?',(pid,),one=True)
    assert (patient['name'],patient['age'],patient['sex'])==('Anaya',25,'Female')
    assert patient['code']=='AUDIO-123'
    assert attach_patient_candidate(db,job_id='123',hospital_id=1,transcript=source,audio_hashes=['sha'])==pid
    with pytest.raises(RuntimeError,match='differs'):
        attach_patient_candidate(db,job_id='123',hospital_id=1,transcript='Mrs. Different age 25 female',audio_hashes=['sha'])
    assert len(db.query('SELECT * FROM patients'))==1


def test_absent_audio_identity_remains_blank_not_fictional(tmp_path):
    db=db_without_patients(tmp_path)
    pid=attach_patient_candidate(db,job_id='no-demographics',hospital_id=1,
           transcript='Appendix microscopy without patient demographics',audio_hashes=['hash'])
    patient=db.query('SELECT * FROM patients WHERE id=?',(pid,),one=True)
    assert patient['name']=='' and patient['age'] is None and patient['sex'] is None
    result,issues,note=verified_audio_patient({'patient_name':'','patient_age':'', 'patient_sex':''},candidate=patient)
    assert issues and note==''


def test_spoken_patient_variants_and_refusal_to_guess():
    assert dictated_patient('Mrs. Anaya age 25 female')=={'name':'Anaya','age':25,'sex':'Female'}
    assert dictated_patient('patient name Ananya age 25 female')=={'name':'Ananya','age':25,'sex':'Female'}
    assert dictated_patient('specimen appendix measuring five centimeters')=={'name':'','age':None,'sex':''}


def test_minor_audio_name_correction_is_audited():
    original={'code':'AUDIO-123','name':'Anaya','age':25,'sex':'Female'}
    altered={'patient_name':'Ananya','patient_age':'25','patient_sex':'Female','confirmed_patient_identity':'yes'}
    _,issues,_=verified_audio_patient(altered,candidate=original)
    assert issues and 'Explain any missing' in issues[0]
    altered['identity_correction_note']='I listened to the original recording and verified the intended spelling.'
    out,issues,note=verified_audio_patient(altered,candidate=original)
    assert not issues and out['name']=='Ananya' and note


def render_page(audio_only):
    root=Path(__file__).resolve().parents[1]
    env=Environment(loader=FileSystemLoader(str(root/'templates')))
    env.globals.update(url_for=lambda *a,**k:'/path',csrf_token=lambda:'secret',
            audio_only_mode=lambda:audio_only,get_flashed_messages=lambda:[])
    return env.get_template('new_report.html').render(who=SimpleNamespace(name='Doctor',role='doctor'),
          request=SimpleNamespace(endpoint='new_report'), audio_only=audio_only, patients=[],
          asr={'provider':'sarvam','model':'saaras:v4','mode':'codemix','ready':True})


def test_audio_mode_form_has_no_patient_selection_or_registry():
    page=render_page(True)
    assert 'No patient selection required' in page
    assert 'name="patient_id"' not in page
    assert 'Choose a synthetic patient' not in page
    assert 'SYNTHETIC PATIENT REGISTRY' not in page
    assert 'name="audio_notes"' in page
    assert 'name="transcript"' not in page
    assert 'AUDIO EVALUATION' in page


def test_legacy_demo_mode_retains_its_separate_optional_form():
    page=render_page(False)
    assert 'name="patient_id"' in page
    assert 'SYNTHETIC PATIENT REGISTRY' in page


def test_docx_labels_asr_patient_unverified(tmp_path):
    patient={'code':'AUDIO-123','name':'Anaya','age':25,'sex':'Female'}
    hospital={'logo_text':'Example Local Evaluation','address':'Nonclinical test'}
    path=tmp_path/'out.docx'
    render_docx({'id':'123456789999'},patient,hospital,[],[],[],[],path,
                assessment={'spoken_patient':{'name':'Anaya','age':25,'sex':'Female'},
                 'blocking_issues':[], 'sections_unverified':{}, 'dates_as_spoken':{}})
    doc=Document(path)
    text=' '.join(x.text for x in doc.paragraphs)+' '+' '.join(c.text for row in doc.tables[0].rows for c in row.cells)
    assert 'Patient name (ASR; unverified)' in text and 'Anaya' in text
    assert 'Selected patient name' not in text
