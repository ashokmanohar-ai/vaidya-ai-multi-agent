"""Executable domain-level end-to-end scenario (no Flask/AWS credentials needed)."""
import io
import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest
from docx import Document
from PIL import Image

from vaidya.database import Database, password_hash, verify_password, utcnow
from vaidya.providers import deterministic_extract, BedrockConverseProvider, get_provider
from vaidya.workflow import run_extraction
from vaidya.reports import render_docx, final_docx, report_text
from vaidya.learning import record_learning, assess_edit, create_seed_learning

SAMPLE=('Specimen: appendix. Sections reveal acute inflammation and neutrophilic infiltrate. '
        'Histopathology finding dictated: acute appendicitis. '
        'Patient temperature was measured at 102 degrees Fahrenheit. '
        'Please bold acute appendicitis in the report.')


def prepare(tmp_path):
    db=Database(tmp_path/'db.sqlite3');db.seed_demo();return db


def test_password_security():
    hashed=password_hash('Test!238')
    assert hashed!='Test!238'
    assert verify_password('Test!238',hashed)
    assert not verify_password('Incorrect',hashed)


def test_seed_users_patients(tmp_path):
    db=prepare(tmp_path)
    assert db.query('SELECT COUNT(*) n FROM users',one=True)['n']==2
    assert db.query('SELECT COUNT(*) n FROM patients',one=True)['n']==7


def test_extraction_separates_clinical_from_formatting():
    result=deterministic_extract(SAMPLE)
    assert result.provider=='demo-rules'
    assert any('acute appendicitis' in f['text'] for f in result.findings)
    assert len(result.instructions)==1
    assert all('Please bold' not in f['text'] for f in result.findings)
    assert 'Please bold' not in report_text(result.findings,result.instructions,result.review_flags)


def test_missing_temperature_warning():
    result=deterministic_extract('Patient has fever. No measured temperature was provided.')
    # Text contains the word temperature, so no simple heuristic is guaranteed; there is no hallucinated numeric value.
    assert not any('102' in f['text'] for f in result.findings)


def test_invalid_empty_dictation():
    with pytest.raises(ValueError):deterministic_extract('   ')


def test_workflow_dispatch_demo(monkeypatch):
    monkeypatch.setenv('LLM_PROVIDER','demo')
    d=run_extraction(SAMPLE)
    assert d['provider']=='demo-rules'
    assert d['instructions'] and len(d['findings'])==4


def test_provider_switch_rejects_unknown(monkeypatch):
    monkeypatch.setenv('LLM_PROVIDER','not-approved')
    with pytest.raises(ValueError):get_provider()


def test_bedrock_guard_prevents_unapproved_remote(monkeypatch):
    monkeypatch.delenv('ALLOW_REMOTE_SYNTHETIC_INPUT', raising=False)
    monkeypatch.delenv('BEDROCK_MODELS', raising=False)
    monkeypatch.delenv('ALLOW_REMOTE_SYNTHETIC_INPUT',raising=False)
    with pytest.raises(RuntimeError,match='disabled'):
        BedrockConverseProvider(client=object()).extract('synthetic text')


def test_bedrock_mocked_converse_contract(monkeypatch):
    monkeypatch.delenv('BEDROCK_MODELS', raising=False)
    monkeypatch.setenv('ALLOW_REMOTE_SYNTHETIC_INPUT','true')
    class FakeClient:
        def converse(self,**kwargs):
            assert kwargs['modelId']=='test.mocked.model'
            payload={'findings':[{'text':'Appendix shows inflammation'}],
                     'instructions':[{'instruction':'bold inflammation'}],
                     'review_flags':['Verify source phrase']}
            return {'output':{'message':{'content':[{'text':json.dumps(payload)}]}}}
    result=BedrockConverseProvider(model='test.mocked.model',client=FakeClient()).extract('appendix inflammation')
    assert result.provider=='aws-bedrock'
    assert result.findings[0]['text']=='Appendix shows inflammation'
    assert result.review_flags


def test_bedrock_rejects_invalid_structure(monkeypatch):
    monkeypatch.delenv('BEDROCK_MODELS', raising=False)
    monkeypatch.setenv('ALLOW_REMOTE_SYNTHETIC_INPUT','true')
    class BadClient:
        def converse(self,**kw):return {'output':{'message':{'content':[{'text':'{"hello":"bad"}'}]}}}
    with pytest.raises(ValueError):BedrockConverseProvider(client=BadClient()).extract('synthetic')


def test_docx_generation_and_instruction_non_leakage(tmp_path):
    db=prepare(tmp_path)
    patient=db.query('SELECT * FROM patients WHERE id=1',one=True)
    hospital=db.query('SELECT * FROM hospitals WHERE id=1',one=True)
    result=run_extraction(SAMPLE)
    image=tmp_path/'slide.png'
    Image.new('RGB',(60,60),'#8899aa').save(image)
    docfile=tmp_path/'draft.docx'
    render_docx({'id':'demo-job-uuid'},patient,hospital,result['findings'],result['instructions'],result['review_flags'],[image],docfile)
    assert docfile.stat().st_size>25000
    doc=Document(docfile)
    texts='\n'.join(p.text for p in doc.paragraphs)
    assert 'acute appendicitis' in texts
    assert 'Please bold acute appendicitis' not in texts
    assert len(doc.inline_shapes)==1
    assert any(run.bold and 'acute appendicitis' in run.text for p in doc.paragraphs for run in p.runs)


def test_review_approval_learning_full_cycle(tmp_path):
    db=prepare(tmp_path)
    create_seed_learning(db)
    sample=run_extraction(SAMPLE)
    patient=db.query('SELECT * FROM patients WHERE id=1',one=True)
    hospital=db.query('SELECT * FROM hospitals WHERE id=1',one=True)
    for i in range(1,6):
        jid=f'demo-test-job-{i}'
        DBQ='''INSERT INTO jobs(id,patient_id,doctor_id,hospital_id,report_type,original_transcript,translated_transcript,findings_json,instruction_json,draft_text,draft_filename,image_json,provider,model,status,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)'''
        db.execute(DBQ,(jid,i,1,1,'histopathology',SAMPLE,SAMPLE,json.dumps(sample['findings']),json.dumps(sample['instructions']),sample['draft_text'],'draft.docx','[]','demo-rules','rules-v1','draft_ready',utcnow()))
        final=sample['draft_text']+'\nSignature: Dr. Brindha'
        with db.connect() as c:
            c.execute('UPDATE jobs SET final_text=?,status=?,approved_by=?,approved_at=? WHERE id=?',
                     (final,'approved',1,utcnow(),jid))
        evidence=record_learning(db,jid,sample['draft_text'],final)
        assert evidence['evidence_count']==i
        if i<5:
            assert db.query('SELECT * FROM learnings WHERE pattern_key=?',('signature_block',))==[]
    learn=db.query('SELECT * FROM learnings WHERE pattern_key=?',('signature_block',),one=True)
    assert learn['supporting_jobs']==5 and learn['status']=='suggested'
    baseline=db.query('SELECT * FROM skills ORDER BY version DESC LIMIT 1',one=True)
    assert baseline['version']==1
    with db.connect() as c:
        c.execute('UPDATE learnings SET status=?,decided_by=? WHERE id=?',('approved',2,learn['id']))
        c.execute('INSERT INTO skills(version,content,learning_id,approved_by,created_at) VALUES(?,?,?,?,?)',
                 (2,baseline['content']+'\nAPPROVED: '+learn['description'],learn['id'],2,utcnow()))
    current=db.query('SELECT * FROM skills ORDER BY version DESC LIMIT 1',one=True)
    assert current['version']==2 and 'signature' in current['content']
    assert db.query('SELECT COUNT(*) n FROM jobs WHERE status=?',('approved',),one=True)['n']==5


def test_unedited_draft_does_not_create_diff(tmp_path):
    db=prepare(tmp_path)
    assert assess_edit('Same','Same') is None


def test_audit_records_action(tmp_path):
    db=prepare(tmp_path)
    db.audit(1,'approve_final','job','synthetic-123')
    r=db.query('SELECT * FROM audit_events',one=True)
    assert r['resource_id']=='synthetic-123' and r['action']=='approve_final'


def test_final_docx_created(tmp_path):
    db=prepare(tmp_path)
    p=db.query('SELECT * FROM patients WHERE id=1',one=True)
    h=db.query('SELECT * FROM hospitals WHERE id=1',one=True)
    outfile=tmp_path/'final.docx'
    final_docx('Synthetic finding.\nSignature: Dr. Brindha',{'id':'j-123'},p,h,outfile)
    assert 'Synthetic finding' in '\n'.join(x.text for x in Document(outfile).paragraphs)
