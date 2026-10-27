"""Synthetic safety regression suite, derived from earlier failure *patterns*.

No identifiable patient recordings or documents are included in the repository.
"""
from io import BytesIO
from pathlib import Path
import json
import zipfile

from docx import Document
from vaidya.clinical_review import assess_report, final_document_issues
from vaidya.database import Database
from vaidya.reports import render_docx, report_text

DEMO_PATIENT = {'name': 'Patient Alpha', 'age': 45, 'sex': 'Female', 'code': 'DEMO-101'}


def assessment(source, findings=None, *, audio=True, patient=None):
    return assess_report(patient or DEMO_PATIENT, source,
                         findings if findings is not None else [{'text': source}], [], audio_present=audio)


def test_name_age_mismatch_is_blocker_and_headers_not_overwritten():
    a=assessment('Patient name Mrs. Sample age 25 female appendix specimen gross measuring 5 cm')
    assert len(a['blocking_issues']) == 2
    assert a['linked_patient']['name'] == 'Patient Alpha'
    assert a['spoken_patient']['name'] == 'Sample'
    assert a['spoken_patient']['age'] == 25
    assert 'Sample' not in a['linked_patient']['name']
    assert final_document_issues('Corrected draft', a)


def test_matching_synthetic_identity_is_not_blocked():
    a=assessment('Patient name Mrs. Patient Alpha age 45 female specimen appendix gross appendix 5 cm', audio=False)
    assert a['blocking_issues'] == []


def test_malformed_dates_remain_unverified_not_rewritten():
    a=assessment('Received 36 reported 47 biopsy number 2007 specimen appendix')
    assert a['dates_as_spoken'] == {'received':'36','reported':'47'}
    assert a['biopsy_id_as_spoken']=='2007'
    assert any('DATE NOT VERIFIED' in x for x in a['review_warnings'])
    assert any('BIOPSY' in x for x in a['review_warnings'])
    assert any('received 36' in x for x in final_document_issues('Received 36 reported 47',a))
    assert final_document_issues('Received 30/06, reported 04/07',a)==[]


def test_negation_inversion_blocks_approval():
    a=assessment('Microscopy shows no atypical cells', findings=[{'text':'Microscopy shows and atypical cells'}])
    assert any('NEGATION' in x for x in a['blocking_issues'])


def test_ambiguous_asr_needs_audio_review_and_corrected_final_text():
    a=assessment('Microscopy appendix and atypical cells. Impression acute inflammation')
    assert any('HIGH RISK ASR' in x for x in a['review_warnings'])
    assert final_document_issues('Microscopy appendix and atypical cells',a)
    assert final_document_issues('Microscopy no atypical cells seen',a)==[]


def test_referral_edit_after_impression_separated():
    sample = ('Patient name Mrs. Patient Alpha age 45 female specimen appendix '
              'gross appendix is 5 cm microscopy shows neutrophils '
              'impression features of inflammation. It is referred by Dr. Doctor One/Dr. Doctor Two.')
    a=assessment(sample)
    assert 'Dr. Doctor One' in ' '.join(a['instruction_candidates'])
    assert 'Doctor One' not in a['sections_unverified']['impression']
    assert 'inflammation' in a['sections_unverified']['impression']


def test_code_mixed_tamil_text_preserved_unmodified():
    sample='Specimen appendix. gross 5 cm. microscopy atypical cells இல்லை. impression inflammation. இது டாக்டர் பெயரை மாற்றவும்.'
    a=assessment(sample)
    assert 'இல்லை' in a['sections_unverified']['microscopy']
    assert 'இது' in a['sections_unverified']['impression']  # Do not pretend to translate Tamil


def test_report_text_is_structured_not_blindly_long_bullets():
    source = 'Specimen appendix. gross 5 cm. microscopy chronic inflammation. impression mild inflammation.'
    a=assessment(source, audio=False)
    txt=report_text([{'text':source}], [],a['review_warnings'],assessment=a)
    assert 'GROSS DESCRIPTION — UNVERIFIED' in txt
    assert 'MICROSCOPY — UNVERIFIED' in txt
    assert 'IMPRESSION — UNVERIFIED' in txt
    assert 'Selected patient: Patient Alpha' in txt


def test_docx_renders_blockers_and_sections_without_modifying_source(tmp_path):
    source='Patient name Mrs. Sample age 25 female gross appendix 5 cm microscopy and atypical cells impression inflammation'
    a=assessment(source)
    out=tmp_path/'demo_review.docx'
    hospital={'logo_text':'SYNTHETIC PATHOLOGY','address':'DEMO ONLY'}
    render_docx({'id':'demo-validation-job'}, DEMO_PATIENT, hospital,
                [{'text':source}], [],a['blocking_issues']+a['review_warnings'],[],out,assessment=a)
    assert zipfile.is_zipfile(out)
    doc=Document(out)
    whole='\n'.join(p.text for p in doc.paragraphs)+'\n'+ '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert 'PATIENT IDENTITY CONFLICT' in whole
    assert 'Patient Alpha' in whole
    assert 'Sample' in whole
    assert 'UNVERIFIED' in whole
    assert 'Gross description' in whole
    assert 'Microscopy' in whole
    assert 'Impression' in whole


def test_legacy_sqlite_migration_adds_validation_column(tmp_path):
    db=Database(tmp_path/'vaidya.sqlite3')
    db.init()
    columns=db.query('PRAGMA table_info(jobs)')
    assert 'validation_json' in {x['name'] for x in columns}
    db.init()  # idempotent


def test_second_negated_mention_does_not_hide_ambiguous_positive_clause():
    a=assessment('Microscopy and atypical cells. Impression benign appendix.',
                 findings=[{'text':'Microscopy and atypical cells.'}])
    final='Microscopy: and atypical cells. In another note no atypical cells were mentioned.'
    assert any('ambiguous atypical-cell' in x for x in final_document_issues(final,a))


def test_unresolved_draft_template_and_placeholders_block_finalization():
    a=assessment('Gross appendix 5 cm.',audio=False)
    texts=[
        'HISTOPATHOLOGY — DRAFT / NOT VALIDATED\nGROSS DESCRIPTION — UNVERIFIED',
        'MICROSCOPY — UNVERIFIED\n[Not clearly dictated]',
        'PATIENT RECORD MUST MATCH SOURCE DICTATION\nSelected patient: Demo',
    ]
    for text in texts:
        assert any('draft warnings' in x for x in final_document_issues(text,a))


def test_reviewed_final_does_not_inherit_draft_template_blockers():
    a=assessment('Gross appendix 5 cm, microscopy no atypical cells.',audio=True)
    final='Gross: Appendix 5 cm. Microscopy: no atypical cells seen. Clinician-reviewed synthetic content.'
    assert final_document_issues(final,a)==[]


def test_dockerfile_copies_recursive_requirements_before_install():
    docker=(Path(__file__).resolve().parents[1]/'deployment'/'Dockerfile').read_text()
    assert 'COPY requirements.txt requirements-production.txt requirements-sarvam.txt ./' in docker
    assert docker.index('COPY requirements.txt requirements-production.txt requirements-sarvam.txt ./') < docker.index('RUN pip install')
    assert 'ASR_PROVIDER=sarvam' not in (Path(__file__).resolve().parents[1]/'.env.production.example').read_text().splitlines()[0:10]
