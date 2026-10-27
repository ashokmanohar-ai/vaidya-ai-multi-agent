"""Deterministic DOCX renderer for clinician-reviewed drafts."""
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path
import io
import re


def report_text(findings, instructions, flags, assessment=None):
    """Editable clinician draft; section candidates are never labelled as verified."""
    lines = ['HISTOPATHOLOGY — DRAFT / NOT VALIDATED']
    if assessment:
        lines.append('PATIENT RECORD MUST MATCH SOURCE DICTATION')
        fields = assessment.get('linked_patient') or {}
        lines.append('Selected patient: ' + str(fields.get('name', '')) + ' / ' + str(fields.get('code', '')))
        spoken = assessment.get('spoken_patient') or {}
        if any(spoken.values()):
            lines.append('Dictation patient (UNVERIFIED): ' + str(spoken))
        sections = assessment.get('sections_unverified') or {}
        for key, title in [('context','SPECIMEN / INTAKE CONTEXT'), ('gross','GROSS DESCRIPTION'),
                           ('microscopy','MICROSCOPY'),('impression','IMPRESSION')]:
            lines.extend(['', title+' — UNVERIFIED', sections.get(key,'') or '[Not clearly dictated]'])
        edits = assessment.get('instruction_candidates') or []
        if edits:
            lines.extend(['', 'DOCUMENT EDIT INSTRUCTIONS — DO NOT TREAT AS CLINICAL FINDINGS'])
            lines.extend(['- ' + x for x in edits])
    else:
        lines.append('DICTATED OBSERVATIONS')
        lines.extend('• ' + str(x.get('text','')) for x in findings)
    if flags:
        lines.extend(['', 'REQUIRES DOCTOR REVIEW'])
        lines.extend('• ' + str(f) for f in flags)
    return '\n'.join(lines)


def render_docx(job, patient, hospital, findings, instructions, flags, images, output,
                approved_skills="", assessment=None):
    """Make a visibly UNVERIFIED report with explicit linked/context discrepancies."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.65)
    section.header_distance = Inches(.3)
    header = section.header.paragraphs[0]
    header.text = hospital['logo_text']
    header.style = 'Title'
    sub = doc.add_paragraph(hospital['address']); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading('HISTOPATHOLOGY REPORT', 1)
    warn = doc.add_paragraph('UNVERIFIED AI-ASSISTED DRAFT — NOT FOR CLINICAL USE')
    for run in warn.runs:
        run.bold = True
        run.font.color.rgb = RGBColor(190, 46, 30)
    doc.add_paragraph('Clinician must compare source recording, patient record, all numbers and negations before approval.')
    t = doc.add_table(rows=0, cols=2)
    t.style = 'Light Shading Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    audio_identity=str(patient['code']).startswith('AUDIO-')
    for left, right in [('Internal ASR candidate reference' if audio_identity else 'Selected patient reference', patient['code']),
                        ('Patient name (ASR; unverified)' if audio_identity else 'Selected patient name', patient['name'] or '[not captured]'),
                        ('Age / sex (ASR; unverified)' if audio_identity else 'Selected age / sex',
                         f"{patient['age'] if patient['age'] is not None else '[not captured]'} / {patient['sex'] or '[not captured]'}"),
                        ('Report type', 'Histopathology'),
                        ('Job reference', job['id'][:12]),
                        ('Status', 'DRAFT — DOCTOR REVIEW')]:
        row = t.add_row().cells
        row[0].text = left
        row[1].text = str(right)
    if assessment:
        spoken = assessment.get('spoken_patient') or {}
        if any(value for value in spoken.values()):
            doc.add_heading('Patient details detected in dictation — unverified', 2)
            doc.add_paragraph('Name: {} | Age: {} | Sex: {}'.format(
                spoken.get('name') or '[not captured]',
                spoken.get('age') if spoken.get('age') is not None else '[not captured]',
                spoken.get('sex') or '[not captured]'))
        if assessment.get('blocking_issues'):
            doc.add_heading('STOP — conflicting source details', 2)
            for issue in assessment['blocking_issues']:
                p = doc.add_paragraph(style='List Bullet')
                run=p.add_run(issue);run.bold=True;run.font.color.rgb=RGBColor(190,46,30)
        if assessment.get('biopsy_id_as_spoken'):
            doc.add_paragraph('Biopsy number as dictated (NOT VERIFIED): '+assessment['biopsy_id_as_spoken'])
        if assessment.get('dates_as_spoken'):
            doc.add_paragraph('Dates as dictated (NOT VERIFIED): '+', '.join(
                key+': '+str(value) for key, value in assessment['dates_as_spoken'].items()))
        sections = assessment.get('sections_unverified') or {}
        for field, title in [('context','Dictated context / specimen'), ('gross','Gross description'),
                             ('microscopy','Microscopy'), ('impression','Impression')]:
            doc.add_heading(title + ' — unverified', 2)
            content=sections.get(field, '').strip()
            doc.add_paragraph(content if content else '[Not clearly dictated; clinician must supply or confirm]')
        candidate_instructions=assessment.get('instruction_candidates') or []
        if candidate_instructions:
            doc.add_heading('Document-edit requests — NOT clinical findings', 2)
            for instruction in candidate_instructions:
                doc.add_paragraph(instruction, style='List Bullet')
    else:
        doc.add_heading('Dictated observations — unverified', level=2)
        bold_targets=[]
        for ins in instructions:
            match=re.search(r'\b(?:please )?(?:bold|highlight)\s+(.+?)(?:\s+in\s+the\s+report)?[.!]?$',ins['instruction'],flags=re.I)
            if match: bold_targets.append(match.group(1).strip(' .!'))
        for fact in findings:
            sentence=str(fact['text'])
            paragraph=doc.add_paragraph(style='List Bullet')
            highlighted=False
            for target in bold_targets:
                m=re.search(re.escape(target),sentence,flags=re.I)
                if m:
                    paragraph.add_run(sentence[:m.start()]); paragraph.add_run(sentence[m.start():m.end()]).bold=True
                    paragraph.add_run(sentence[m.end():]); highlighted=True; break
            if not highlighted: paragraph.add_run(sentence)
        if not findings:
            doc.add_paragraph('[No dictated observations extracted; doctor review required]')
    if images:
        doc.add_heading('Uploaded reference images', 2)
        for index, path in enumerate(images,1):
            doc.add_paragraph(f'Image {index} · reference only; no AI diagnosis performed')
            try: doc.add_picture(str(path), width=Inches(3.5))
            except Exception: doc.add_paragraph('[Image could not be embedded]')
    if flags:
        doc.add_heading('Items requiring clinician clarification', 2)
        for flag in list(dict.fromkeys(flags)):
            doc.add_paragraph(str(flag),style='List Bullet')
    if 'APPROVED:' in approved_skills and 'signature' in approved_skills.lower():
        doc.add_paragraph('\nClinician signature: _____________________      Date: __________')
    footer = section.footer.paragraphs[0]
    footer.text = 'VAIDYA AI · SYNTHETIC DEVELOPMENT DEMO · NOT A VALIDATED MEDICAL DOCUMENT'
    Path(output).parent.mkdir(parents=True,exist_ok=True)
    doc.save(str(output))
    return str(output)


def final_docx(text,job,patient,hospital,output):
    doc=Document()
    doc.add_heading(hospital['logo_text'],0)
    doc.add_paragraph('CLINICIAN-APPROVED SYNTHETIC DEMONSTRATION')
    doc.add_heading('Histopathology Report',1)
    doc.add_paragraph(f"Patient: {patient['name']} ({patient['code']})")
    for l in text.splitlines():
        if l.strip() in ("HISTOPATHOLOGY — DRAFT / NOT VALIDATED", "DICTATED OBSERVATIONS"): continue
        doc.add_paragraph(l)
    doc.add_paragraph(f"Approved by demo clinician · job {job['id'][:12]}")
    doc.save(str(output))


def structured_final_docx(fields, job, patient, hospital, output):
    """Render only manually verified structured fields. Synthetic-only pilot."""
    from .structured_review import FIELDS
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(.62);sec.bottom_margin=Inches(.58)
    sec.left_margin=Inches(.78);sec.right_margin=Inches(.78)
    header=sec.header.paragraphs[0]
    header.text=hospital['logo_text'];header.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:run.bold=True;run.font.size=Pt(16);run.font.color.rgb=RGBColor(25,65,94)
    addr=doc.add_paragraph(hospital.get('address',''))
    addr.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading('HISTOPATHOLOGY REPORT',1)
    notice=doc.add_paragraph('REVIEWED AUDIO-DERIVED FIELDS · DEVELOPMENT EVALUATION · NOT FOR CLINICAL USE')
    for run in notice.runs:run.font.size=Pt(8);run.bold=True;run.font.color.rgb=RGBColor(150,40,36)
    tbl=doc.add_table(rows=0,cols=2);tbl.style='Light Shading Accent 1'
    detail=[('Patient reference',patient['code']),('Patient name',patient['name']),
            ('Age / Sex',f"{patient['age']} / {patient['sex']}"),
            ('Referring doctor(s)',fields['referring_doctors']),
            ('Specimen',fields['specimen']),('Biopsy / accession number',fields['biopsy_id']),
            ('Received',fields['received_date']),('Reported',fields['reported_date'])]
    for label,value in detail:
        cells=tbl.add_row().cells
        cells[0].text=label;cells[1].text=str(value)
        for cell in cells:cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for key,title in (('gross','GROSS DESCRIPTION'),('microscopy','MICROSCOPY'),('impression','IMPRESSION')):
        doc.add_heading(title,2)
        doc.add_paragraph(fields[key])
    doc.add_paragraph('Reviewed by authorised evaluator; all values require independent verification against source evidence.')
    doc.add_paragraph(f"Report audit reference: {job['id'][:12]}")
    footer=sec.footer.paragraphs[0]
    footer.text='VAIDYA AI · DEVELOPMENT EVALUATION · NOT A VALIDATED MEDICAL DOCUMENT'
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    Path(output).parent.mkdir(parents=True,exist_ok=True)
    doc.save(str(output))
    return str(output)
